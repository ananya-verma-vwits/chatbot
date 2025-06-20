import os
from fastapi import UploadFile, HTTPException
from PyPDF2 import PdfReader
import docx
import fitz  # PyMuPDF
import camelot.io as camelot

FILES_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), "files")  # Save files inside app/files
os.makedirs(FILES_DIR, exist_ok=True)

def save_file(file: UploadFile) -> str:
    """Save an uploaded file to the server."""
    file_path = os.path.join(FILES_DIR, str(file.filename))
    try:
        with open(file_path, "wb") as buffer:
            buffer.write(file.file.read())
        return file_path
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error saving file: {str(e)}")

def extract_text_from_pdf(file_path: str) -> str:
    """Extract text from a PDF file."""
    try:
        reader = PdfReader(file_path)
        text = ""
        for page in reader.pages:
            text += page.extract_text()
        return text
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error reading PDF: {str(e)}")

def extract_text_from_docx(file_path: str) -> str:
    """Extract text from a DOCX file."""
    try:
        doc = docx.Document(file_path)
        text = "\n".join([paragraph.text for paragraph in doc.paragraphs])
        return text
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error reading DOCX: {str(e)}")
    


def extract_images_from_pdf(file_path: str) -> list:
    """Extract images from a PDF file."""
    try:
        images = []
        pdf_document = fitz.open(file_path)
        for page_number in range(len(pdf_document)):
            page = pdf_document[page_number]
            image_list = page.get_images(full=True)
            for img_index, img in enumerate(image_list):
                xref = img[0]
                base_image = pdf_document.extract_image(xref)
                image_bytes = base_image["image"]
                images.append(image_bytes)
        return images
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error extracting images: {str(e)}")
    


def extract_tables_from_pdf(file_path: str) -> list:
    """Extract tables from a PDF file."""
    try:
        tables = camelot.read_pdf(file_path, pages="all")
        extracted_tables = [table.df for table in tables]
        return extracted_tables
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error extracting tables: {str(e)}")
    
def process_pdf(file_path: str):
    """Process a PDF file to extract text, images, and tables."""
    text = extract_text_from_pdf(file_path)
    images = extract_images_from_pdf(file_path)
    tables = extract_tables_from_pdf(file_path)

    # Combine extracted content
    context = {
        "text": text,
        "images": images,
        "tables": tables
    }
    return context

def process_pdf_to_markdown(file_path: str) -> str:
    """Convert a PDF file to Markdown format."""
    context = process_pdf(file_path)  # Extract text, images, and tables

    markdown_content = f"# Document Content\n\n{context['text']}\n\n"

    # Add tables to Markdown
    markdown_content += "## Tables\n\n"
    for i, table in enumerate(context['tables']):
        markdown_content += f"### Table {i + 1}\n\n"
        markdown_content += table.to_markdown(index=False) + "\n\n"

    # Add image placeholders to Markdown
    markdown_content += "## Images\n\n"
    for i, image in enumerate(context['images']):
        markdown_content += f"![Image {i + 1}](image_{i + 1}.png)\n\n"

    return markdown_content